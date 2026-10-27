#!/usr/bin/env python3
"""AI visibility check (v16.0): 7 questions x 10 runs x 7 tools on Oct 27 2026, 5 tools from Oct 28 2026 (v16.1), row 9, the v16 listing-site rules,
the spreadsheet. Rounds: argv[1] (default 3)."""
import copy, json, pathlib, sys, tempfile
ENG = pathlib.Path(__file__).resolve().parent.parent; sys.path.insert(0, str(ENG))
import ai_visibility as V, ai_answers as AA, score as S, alert_cards as AC, directories as D

EX = json.loads((ENG / "examples/evidence-example-v16.json").read_text())
fails = 0
def ok(name, cond, extra=""):
    global fails
    print(("PASS " if cond else "FAIL ") + name + ("" if cond else f"  -> {extra}")); fails += 0 if cond else 1

def run_once():
    ev = copy.deepcopy(EX)
    ok("the worked example is complete", not S.check(copy.deepcopy(ev)), S.check(copy.deepcopy(ev)))
    ok("490 answers are expected on Oct 27 2026", len(ev["ai_visibility"]["answers"]) == len(V.tools(ev)) * len(V.QUESTIONS) * V.RUNS == 490)
    ok("no question names a dealer or a brand", all("{city_st}" in q and "Mitsubishi" not in q for _, _, q in V.QUESTIONS))
    ok("3 Sales, 2 Service, 2 Parts", [d for _, d, _ in V.QUESTIONS] == ["sales"] * 3 + ["service"] * 2 + ["parts"] * 2)
    # scoring
    sc = V.scores(ev); A = [a for a in ev["ai_visibility"]["answers"] if a["status"] == "answered"]
    ok("blocked answers are not counted", sc["answered"] == len(A) == 476, sc["answered"])
    ok("row 9 = named / answered x 10", abs(sc["row"] - 10 * sum(a["named"] for a in A) / len(A)) < 1e-9)
    q = sc["questions"]["sales_dealership"]; ok("a question scores named / answered x 100", q["score"] == round(100 * q["named"] / q["answered"]), q)
    raw, rows = S.score_rows(copy.deepcopy(ev))
    ok("v16 audits score nine rows", len(rows) == 9 and S.row_names(rows)[8] == "AI Answer Visibility", rows)
    ok("status can be set by row 9", S.status(rows)["lowest_row"] == "AI Answer Visibility", S.status(rows))
    ok("overall is the average of nine", S.overall(rows) == S.r1d(sum(rows) / 9))
    old = copy.deepcopy(ev); old["audit_date"] = "2026-10-26"
    ok("an audit dated Oct 26 2026 keeps eight rows and needs the v15.1 check", len(S.score_rows(copy.deepcopy(old))[1]) == 8 and any("ai_answers" in p for p in S.check(copy.deepcopy(old))))
    ok("the v15.1 check is not required from Oct 27 2026", not AA.check(ev))
    every = copy.deepcopy(ev)
    for a in every["ai_visibility"]["answers"]:
        if a["status"] == "answered": a["named"] = True
    ok("named in every answer -> row 9 is 10", S.score_rows(every)[1][8] == 10.0)
    # the card
    cards, _ = AC.pick(S.resolve_ages(copy.deepcopy(ev)), rows)
    ok("the Alert shows the AI_ANSWERS_NAMED card when row 9 is weakest", cards and cards[0]["head"] == "AI_ANSWERS_NAMED" and cards[0]["target"] == sc["named"], cards[:1])
    ok("stat header fits 18 characters", len("> AI_ANSWERS_NAMED") <= 18)
    # refusals
    def broken(mut, needle, name):
        e = copy.deepcopy(ev); mut(e["ai_visibility"]); p = V.check(e); ok("refuses: " + name, any(needle in x for x in p), p[:2])
    broken(lambda a: a["answers"].pop(0), "not recorded", "a missing answer")
    broken(lambda a: a["answers"].append(copy.deepcopy(a["answers"][1])), "recorded twice", "the same answer twice")
    broken(lambda a: a["answers"][0].update(asked="best Mitsubishi dealer near Vacaville, CA"), ".asked", "a question that adds a brand")
    broken(lambda a: a["answers"][0].update(tool="grok"), "unknown tool", "a tool not on the list")
    broken(lambda a: a["answers"][0].update(run=11), "unknown tool", "an 11th run")
    broken(lambda a: a["answers"][0].update(answer=""), ".answer", "no word-for-word answer")
    broken(lambda a: a["answers"][0].update(businesses=[{"name": "X", "type": "car wash"}]), ".businesses", "an unknown business type")
    broken(lambda a: a["answers"][0].update(named="yes"), ".named", "named not true/false")
    broken(lambda a: a["answers"][0].update(status="blocked", reason=""), ".reason", "blocked with no reason")
    broken(lambda a: a.update(city_st="Vacaville"), "city_st", "no state")
    e = copy.deepcopy(ev); e.pop("ai_visibility"); ok("no AI visibility check -> not scored", any(p.startswith("ai_visibility: missing") for p in S.check(e)))
    # v16 listing-site rules
    R = {r["directory"]: r for r in D.results(copy.deepcopy(ev))}
    ok("no phone shown counts as not agreeing", R["cars_com"]["agree"] is False and "no phone shown" in R["cars_com"]["why"], R["cars_com"])
    ok("a legal company name is skipped, not counted", R["data_axle"]["status"] == "skipped" and R["data_axle"]["agree"] is None, R["data_axle"])
    ok("the closest retailer goes in the rep-brief notes", any("Example Mitsubishi of Fairfield (14 mi)" in n for n in D.notes(ev)), D.notes(ev))
    e = copy.deepcopy(ev); next(d for d in e["nap"]["directories"] if d["directory"] == "autotrader").pop("closest")
    ok("refuses: a car-shopper site with no listing and no closest retailer", any(".closest" in p for p in D.check(e)), D.check(e))
    e = copy.deepcopy(ev); e["audit_date"] = "2026-10-26"; e["ai_answers"] = {"method": "x", "answers": []}
    R = {r["directory"]: r for r in D.results(e)}
    ok("before Oct 27 2026 the legal name is still compared", R["data_axle"]["status"] == "found", R["data_axle"])
    # v16.1: 5 tools through their APIs from Oct 28 2026
    n = copy.deepcopy(ev); n["audit_date"] = "2026-10-28"
    n["ai_visibility"]["answers"] = [a for a in n["ai_visibility"]["answers"] if a["tool"] not in ("copilot", "perplexity")]
    ok("350 answers are expected from Oct 28 2026", len(n["ai_visibility"]["answers"]) == len(V.tools(n)) * len(V.QUESTIONS) * V.RUNS == 350)
    ok("Copilot and Perplexity are not asked from Oct 28 2026", not {"copilot", "perplexity"} & set(V.tools(n)))
    ok("a 5-tool audit dated Oct 28 2026 passes the check", not V.check(copy.deepcopy(n)), V.check(copy.deepcopy(n))[:2])
    e = copy.deepcopy(n); e["ai_visibility"]["answers"][0]["tool"] = "copilot"
    ok("refuses: a Copilot answer on an Oct 28 2026 audit", any("unknown tool" in x for x in V.check(e)), V.check(e)[:2])
    e = copy.deepcopy(ev); e["ai_visibility"]["answers"] = n["ai_visibility"]["answers"]
    ok("an Oct 27 2026 audit still needs all 7 tools", any("140 of 490" in x for x in V.check(e)), V.check(e)[-1:])
    ok("the brief has no Copilot row from Oct 28 2026", "| Copilot |" not in V.brief(n) and "| Copilot |" in V.brief(ev))
    raw, rows = S.score_rows(copy.deepcopy(n)); ok("a 5-tool audit scores row 9", len(rows) == 9, rows)
    # outputs
    W = V.winners(ev); ok("the winner is the business named most", W[("service_oil_change", "chatgpt")][0] == "Example Quick Lube", W[("service_oil_change", "chatgpt")])
    b = V.brief(ev); ok("the brief has one table row per tool per question", b.count("| ChatGPT |") == len(V.QUESTIONS))
    with tempfile.TemporaryDirectory() as t:
        from openpyxl import load_workbook
        p = V.sheet(ev, pathlib.Path(t) / "x.xlsx"); wb = load_workbook(p)
        ok("one spreadsheet tab per tool, named for the tool", wb.sheetnames == [V.NAME[x] for x in V.tools(ev)], wb.sheetnames)
        ok("one row per answer on each tab", all(wb[V.NAME[x]].max_row == 1 + len(V.QUESTIONS) * V.RUNS for x in V.tools(ev)))
    ok("file name", V.file_name("James Mitsubishi Greece", "2026-10-27") == "james-mitsubishi-greece_ai-answers_10-27-26.xlsx")

rounds = int(sys.argv[1]) if len(sys.argv) > 1 else 3
for r in range(rounds): print(f"--- round {r + 1}"); run_once()
print("\nALL PASSED" if not fails else f"\n{fails} FAILED"); sys.exit(1 if fails else 0)
